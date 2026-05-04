"""
NailVesta 中台运营周报生成器 v3 ✨
==================================
作者：为 Chenhao（奶瓶）打造

输入（每周固定上传 3 个文件）：
    1. 28 天订单数据（All_order csv/xlsx）
    2. 28 天退货数据（Return_Refund_Orders xlsx）
    3. 产品图册（NailVesta_产品图册 csv）— 用于供应商和新品识别

时间窗模式（侧边栏可切换）：
    A. 28 天对半切：前 14 天 vs 后 14 天（默认）
    B. 自然周对比：最近完整一周 vs 上一周

新增功能（v3）：
    - 供应商分析：按厂家聚合所有在售款的整体退货率
    - 客诉响应时效（处理周期分桶）
    - 重复退货买家 TOP 10
    - 差评关键词 TOP 10（从 Buyer Note 文本抽取）
    - 包裹丢失分析（Missing package 趋势）
    - 产品生命周期退货率（按上架时长分组）

启动：
    streamlit run app.py
"""

import io
import re
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook


# ╔════════════════════════════════════════════════════════════════════╗
# ║   第 1 部分：常量配置                                              ║
# ╚════════════════════════════════════════════════════════════════════╝

INFLUENCER_RULE_MAP = {
    "SKU Subtotal After Discount = 0": "SKU Subtotal After Discount",
    "SKU Unit Original Price = 0": "SKU Unit Original Price",
    "Order Amount = 0": "Order Amount",
}

# 模板里要复制的基础 sheet
BASE_SHEET = "4.20-4.26"

# 新品阈值
NEW_PRODUCT_DAYS = 30

# 退货文件列名候选
RETURN_COL_CANDIDATES = {
    "qty":              ["Return Quantity", "Sku Quantity of return", "Quantity", "退货数量"],
    "unit_price":       ["Return unit price", "Unit Price", "退款单价"],
    "amount":           ["Order Refund Amount", "Refund Amount", "退款金额"],
    "reason":           ["Return Reason", "Refund Reason", "Reason", "退货原因"],
    "sku_name":         ["SKU Name", "Variation", "SKU Variation", "款式"],
    "sku":              ["Seller SKU", "SKU", "Seller_SKU"],
    "return_status":    ["Return Status", "退货状态"],
    "return_substatus": ["Return Sub Status", "Return SubStatus", "退货子状态"],
    "time_requested":   ["Time Requested", "申请时间"],
    "refund_time":      ["Refund Time", "退款时间"],
    "buyer_username":   ["Buyer Username", "买家用户名"],
    "buyer_note":       ["Buyer Note", "买家备注"],
    "order_id":         ["Order ID", "OrderID", "订单号"],
}

EXCLUDED_SUBSTATUS_KEYWORDS = ["request canceled", "request cancelled"]

CATALOG_COL_CANDIDATES = {
    "sku":          ["SKU", "sku"],
    "style_en":     ["款式英文名称", "英文名称", "Style"],
    "supplier":     ["厂家", "供应商", "Supplier"],
    "list_date":    ["上架时间", "上架日期", "Listed Date"],
    "list_status":  ["上架状态", "状态"],
    "is_listed":    ["是否上架"],
}

# 客诉响应时效分桶（小时）
RESPONSE_TIME_BUCKETS = [
    ("< 6 小时",   0,    6),
    ("6-24 小时",  6,    24),
    ("1-3 天",     24,   72),
    ("3-7 天",     72,   168),
    ("> 7 天",     168,  float("inf")),
]

# 产品生命周期分桶（天）
LIFECYCLE_BUCKETS = [
    ("新品 (≤30天)",       0,    30),
    ("成长期 (31-90天)",   30,   90),
    ("成熟期 (91-180天)",  90,   180),
    ("长尾期 (>180天)",    180,  float("inf")),
]

# 关键词分析的停用词（不计入关键词统计）
STOP_WORDS = {
    "i", "me", "my", "the", "a", "an", "is", "it", "to", "for", "of", "in",
    "on", "at", "and", "or", "but", "with", "this", "that", "they", "them",
    "was", "were", "are", "be", "been", "have", "has", "had", "do", "does",
    "did", "will", "would", "should", "can", "could", "no", "not", "so",
    "as", "if", "by", "from", "your", "you", "we", "us", "our", "just",
    "got", "get", "didn", "don", "wasn", "isn", "haven", "hasn", "won",
    "really", "very", "too", "much", "more", "all", "some", "any", "only",
    "still", "even", "also", "now", "then", "than", "when", "where", "why",
    "how", "what", "who", "which", "out", "up", "down", "off", "back",
    "again", "one", "two", "first", "second", "last", "next", "many",
    "good", "bad", "sure", "well", "fine", "nice", "ok", "okay", "yes",
    "thanks", "please", "hello", "hi", "thank",
}

# 关键词分析中的「重要词组」—— 这些词组出现率远比单词更有信号
IMPORTANT_PHRASES = [
    "wrong size", "doesn't fit", "doesnt fit", "don't fit", "dont fit",
    "too small", "too big", "too short", "too long", "too tight", "too loose",
    "missing", "broken", "defective", "wrong color", "wrong item",
    "fell off", "didn't arrive", "didnt arrive", "never arrived",
    "lost package", "package lost", "wrong shape", "wrong design",
    "poor quality", "bad quality", "low quality", "uneven", "rough",
    "thumb", "pinky", "size chart", "sizing", "fit my", "match",
]


# ╔════════════════════════════════════════════════════════════════════╗
# ║   第 2 部分：通用工具                                              ║
# ╚════════════════════════════════════════════════════════════════════╝

def _read_any(file) -> pd.DataFrame:
    if file is None:
        return pd.DataFrame()
    name = file.name.lower()
    file.seek(0)
    if name.endswith(".csv"):
        return pd.read_csv(file)
    return pd.read_excel(file)


def _find_col(df: pd.DataFrame, candidates: list) -> str:
    for c in candidates:
        if c in df.columns:
            return c
    return None


def extract_style(variation) -> str:
    if pd.isna(variation):
        return ""
    return str(variation).split(",")[0].strip()


def extract_sku_root(sku) -> str:
    if pd.isna(sku):
        return ""
    return str(sku).split("-")[0].strip()


def _bucket_sku(n: int) -> str:
    if n == 1: return "1"
    if n == 2: return "2"
    if n == 3: return "3"
    if n == 4: return "4"
    return "4+"


def _wow(this: float, last: float) -> float:
    if not last:
        return 0
    return this / last - 1


def _parse_pct(s: str):
    if not s:
        return None
    try:
        v = float(s.replace("%", "").strip())
        return v / 100
    except ValueError:
        return None


def _is_request_canceled(substatus_value) -> bool:
    if pd.isna(substatus_value):
        return False
    s = str(substatus_value).lower()
    return any(kw in s for kw in EXCLUDED_SUBSTATUS_KEYWORDS)


def _clean_tab_str(s: pd.Series) -> pd.Series:
    return s.astype(str).str.strip().str.rstrip("\t").str.strip()


# ╔════════════════════════════════════════════════════════════════════╗
# ║   第 3 部分：时间窗拆分                                            ║
# ╚════════════════════════════════════════════════════════════════════╝

def split_time_windows(orders_df: pd.DataFrame, returns_df: pd.DataFrame,
                       mode: str) -> dict:
    """将 28 天数据拆为「上周/本周」"""
    if orders_df.empty:
        raise ValueError("订单数据为空")

    o = orders_df.copy()
    if "Created Time" in o.columns:
        o["__t"] = pd.to_datetime(_clean_tab_str(o["Created Time"]), errors="coerce")
    else:
        raise ValueError("订单文件缺少 Created Time 列")

    r = returns_df.copy() if not returns_df.empty else pd.DataFrame()
    if not r.empty:
        time_col = _find_col(r, RETURN_COL_CANDIDATES["time_requested"])
        if time_col:
            cleaned = _clean_tab_str(r[time_col])
            r["__t"] = pd.to_datetime(cleaned, format="%d/%m/%Y %H:%M:%S", errors="coerce")
            mask_na = r["__t"].isna()
            if mask_na.any():
                r.loc[mask_na, "__t"] = pd.to_datetime(cleaned[mask_na], errors="coerce")
        else:
            r["__t"] = pd.NaT

    t_min = o["__t"].min()
    t_max = o["__t"].max()
    if pd.isna(t_min) or pd.isna(t_max):
        raise ValueError("订单时间字段无法解析")

    if mode == "halve":
        midpoint = t_min + (t_max - t_min) / 2
        last_start, last_end = t_min, midpoint
        this_start, this_end = midpoint, t_max + timedelta(seconds=1)
        last_label = f"{last_start.strftime('%-m.%-d')}-{(last_end - timedelta(days=1)).strftime('%-m.%-d')}"
        this_label = f"{this_start.strftime('%-m.%-d')}-{t_max.strftime('%-m.%-d')}"

    elif mode == "weekly":
        if t_max.weekday() != 6:
            this_end = t_max - timedelta(days=t_max.weekday() + 1)
            this_end = this_end.replace(hour=23, minute=59, second=59)
        else:
            this_end = t_max
        this_end = pd.Timestamp(this_end.replace(hour=23, minute=59, second=59))
        this_start = (this_end - timedelta(days=6)).replace(hour=0, minute=0, second=0)
        last_end = (this_start - timedelta(seconds=1))
        last_start = (last_end - timedelta(days=6)).replace(hour=0, minute=0, second=0)
        last_label = f"{last_start.strftime('%-m.%-d')}-{last_end.strftime('%-m.%-d')}"
        this_label = f"{this_start.strftime('%-m.%-d')}-{this_end.strftime('%-m.%-d')}"
    else:
        raise ValueError(f"未知时间窗模式: {mode}")

    last_orders = o[(o["__t"] >= last_start) & (o["__t"] <= last_end)].copy()
    this_orders = o[(o["__t"] >= this_start) & (o["__t"] <= this_end)].copy()

    if not r.empty and "__t" in r.columns:
        last_returns = r[(r["__t"] >= last_start) & (r["__t"] <= last_end)].copy()
        this_returns = r[(r["__t"] >= this_start) & (r["__t"] <= this_end)].copy()
    else:
        last_returns = pd.DataFrame()
        this_returns = pd.DataFrame()

    return {
        "last_orders": last_orders, "this_orders": this_orders,
        "last_returns": last_returns, "this_returns": this_returns,
        "last_label": last_label, "this_label": this_label,
        "ref_date": t_max, "t_min": t_min, "t_max": t_max,
    }


# ╔════════════════════════════════════════════════════════════════════╗
# ║   第 4 部分：产品图册                                              ║
# ╚════════════════════════════════════════════════════════════════════╝

def process_catalog(df: pd.DataFrame) -> pd.DataFrame:
    """清洗图册，输出标准格式：sku_root | style_en | supplier | list_date | is_active"""
    if df.empty:
        return pd.DataFrame(columns=["sku_root", "style_en", "supplier", "list_date", "is_active"])

    sku_col      = _find_col(df, CATALOG_COL_CANDIDATES["sku"])
    style_col    = _find_col(df, CATALOG_COL_CANDIDATES["style_en"])
    supplier_col = _find_col(df, CATALOG_COL_CANDIDATES["supplier"])
    list_col     = _find_col(df, CATALOG_COL_CANDIDATES["list_date"])
    is_listed_col= _find_col(df, CATALOG_COL_CANDIDATES["is_listed"])

    out = pd.DataFrame()
    out["sku_root"]  = df[sku_col].astype(str).str.strip() if sku_col else ""
    out["style_en"]  = df[style_col].astype(str).str.strip() if style_col else ""
    out["supplier"]  = df[supplier_col].astype(str).str.strip() if supplier_col else ""
    out["list_date"] = pd.to_datetime(df[list_col], errors="coerce") if list_col else pd.NaT
    if is_listed_col:
        out["is_active"] = pd.to_numeric(df[is_listed_col], errors="coerce").fillna(0) > 0
    else:
        out["is_active"] = True
    out["supplier"] = out["supplier"].replace({"nan": "", "NaN": "", "None": ""})
    return out


# ╔════════════════════════════════════════════════════════════════════╗
# ║   第 5 部分：订单数据处理                                          ║
# ╚════════════════════════════════════════════════════════════════════╝

def _empty_orders() -> dict:
    return {
        "orders": 0, "sku_sold": 0, "gmv": 0,
        "aov": 0, "asp": 0, "attach_rate": 0,
        "influencer_orders": 0,
        "bucket_struct": pd.DataFrame(columns=["bucket", "order_count"]),
        "bucket_aov": {},
        "style_sales": pd.DataFrame(columns=["style", "sales", "sku"]),
        "raw_paid": pd.DataFrame(),
    }


def process_orders(df: pd.DataFrame, influencer_col: str) -> dict:
    if df.empty:
        return _empty_orders()

    df = df.copy()
    for col in ["SKU Subtotal After Discount", "SKU Unit Original Price",
                "Order Amount", "Quantity"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)

    df["is_influencer"] = df[influencer_col] == 0

    if "Order Status" in df.columns:
        df_valid = df[df["Order Status"] != "Canceled"].copy()
    else:
        df_valid = df.copy()

    paid = df_valid[~df_valid["is_influencer"]].copy()
    influencer = df_valid[df_valid["is_influencer"]].copy()

    paid["style"] = paid["Variation"].apply(extract_style)
    paid["sku_root"] = paid["Seller SKU"].apply(extract_sku_root)
    influencer["style"] = influencer["Variation"].apply(extract_style)

    paid_orders = paid["Order ID"].nunique()
    paid_sku_sold = int(paid["Quantity"].sum())
    paid_gmv = float(paid["SKU Subtotal After Discount"].sum())

    aov = paid_gmv / paid_orders if paid_orders else 0
    asp = paid_gmv / paid_sku_sold if paid_sku_sold else 0
    attach_rate = paid_sku_sold / paid_orders if paid_orders else 0

    sku_per_order = paid.groupby("Order ID")["Quantity"].sum().rename("sku_count").reset_index()
    sku_per_order["bucket"] = sku_per_order["sku_count"].apply(_bucket_sku)
    bucket_struct = sku_per_order.groupby("bucket").agg(order_count=("Order ID", "count")).reset_index()

    order_amt = paid.groupby("Order ID")["SKU Subtotal After Discount"].sum().rename("order_amt").reset_index()
    sku_per_order = sku_per_order.merge(order_amt, on="Order ID", how="left")
    bucket_aov = sku_per_order.groupby("bucket")["order_amt"].mean().to_dict()

    style_sales = (
        paid.groupby("style")
        .agg(
            sales=("Quantity", "sum"),
            sku=("sku_root", lambda s: s.mode().iloc[0] if not s.mode().empty else ""),
        )
        .reset_index()
        .sort_values("sales", ascending=False)
    )

    return {
        "orders": paid_orders, "sku_sold": paid_sku_sold, "gmv": paid_gmv,
        "aov": aov, "asp": asp, "attach_rate": attach_rate,
        "influencer_orders": influencer["Order ID"].nunique(),
        "bucket_struct": bucket_struct, "bucket_aov": bucket_aov,
        "style_sales": style_sales, "raw_paid": paid,
    }


# ╔════════════════════════════════════════════════════════════════════╗
# ║   第 6 部分：退货数据处理                                          ║
# ╚════════════════════════════════════════════════════════════════════╝

def _empty_returns() -> dict:
    return {
        "total_qty": 0, "total_amount": 0,
        "by_style": pd.DataFrame(columns=["style", "return_qty"]),
        "by_reason": pd.DataFrame(columns=["reason", "qty"]),
        "by_reason_top_style": {},
        "requested_qty": 0, "request_canceled_qty": 0,
        "raw_actual": pd.DataFrame(),  # 实际退货 df，用于后续 KPI 分析
        "raw_all": pd.DataFrame(),     # 全量退货 df（含撤销）
    }


def process_returns(df: pd.DataFrame, exclude_request_canceled: bool = True) -> dict:
    if df.empty:
        return _empty_returns()

    df = df.copy()
    qty_col       = _find_col(df, RETURN_COL_CANDIDATES["qty"])
    unit_price_col= _find_col(df, RETURN_COL_CANDIDATES["unit_price"])
    amt_col       = _find_col(df, RETURN_COL_CANDIDATES["amount"])
    reason_col    = _find_col(df, RETURN_COL_CANDIDATES["reason"])
    sku_name_col  = _find_col(df, RETURN_COL_CANDIDATES["sku_name"])
    sku_col       = _find_col(df, RETURN_COL_CANDIDATES["sku"])
    substatus_col = _find_col(df, RETURN_COL_CANDIDATES["return_substatus"])

    if qty_col:
        df[qty_col] = pd.to_numeric(df[qty_col], errors="coerce").fillna(0)
        df = df[df[qty_col] > 0].copy()
    else:
        df["__qty"] = 1
        qty_col = "__qty"

    if sku_name_col:
        df["style"] = df[sku_name_col].apply(extract_style)
    elif sku_col:
        df["style"] = df[sku_col].apply(extract_sku_root)
    else:
        df["style"] = ""

    # SKU 根（用于后面关联图册）
    if sku_col:
        df["sku_root"] = df[sku_col].apply(extract_sku_root)
    else:
        df["sku_root"] = ""

    requested_qty = int(df[qty_col].sum())

    if substatus_col:
        df["__is_req_canceled"] = df[substatus_col].apply(_is_request_canceled)
    else:
        df["__is_req_canceled"] = False

    request_canceled_qty = int(df[df["__is_req_canceled"]][qty_col].sum())

    df_actual = df[~df["__is_req_canceled"]].copy() if exclude_request_canceled else df.copy()
    total_qty = int(df_actual[qty_col].sum()) if not df_actual.empty else 0

    if not df_actual.empty:
        if unit_price_col:
            df_actual[unit_price_col] = pd.to_numeric(df_actual[unit_price_col], errors="coerce").fillna(0)
            total_amount = float((df_actual[unit_price_col] * df_actual[qty_col]).sum())
        elif amt_col:
            df_actual[amt_col] = pd.to_numeric(df_actual[amt_col], errors="coerce").fillna(0)
            total_amount = float(df_actual[amt_col].sum())
        else:
            total_amount = 0
    else:
        total_amount = 0

    if not df_actual.empty:
        by_style = (
            df_actual.groupby("style")[qty_col].sum().rename("return_qty").reset_index()
            .sort_values("return_qty", ascending=False)
        )
    else:
        by_style = pd.DataFrame(columns=["style", "return_qty"])

    if reason_col:
        by_reason = (
            df.groupby(reason_col)[qty_col].sum().rename("qty").reset_index()
            .rename(columns={reason_col: "reason"})
            .sort_values("qty", ascending=False)
        )
        by_reason_top = {}
        for r, sub in df.groupby(reason_col):
            top = sub.groupby("style")[qty_col].sum().sort_values(ascending=False).head(2)
            by_reason_top[r] = "; ".join([f"{s} ({int(q)})" for s, q in top.items() if s])
    else:
        by_reason = pd.DataFrame(columns=["reason", "qty"])
        by_reason_top = {}

    return {
        "total_qty": total_qty, "total_amount": total_amount,
        "by_style": by_style, "by_reason": by_reason,
        "by_reason_top_style": by_reason_top,
        "requested_qty": requested_qty,
        "request_canceled_qty": request_canceled_qty,
        "raw_actual": df_actual,
        "raw_all": df,
    }


# ╔════════════════════════════════════════════════════════════════════╗
# ║   第 7 部分：新品 + 供应商分析                                     ║
# ╚════════════════════════════════════════════════════════════════════╝

def build_new_products_top10(this_orders: dict, this_returns: dict,
                              catalog: pd.DataFrame, ref_date: datetime) -> pd.DataFrame:
    """近 30 天新品 TOP 10。"""
    if catalog.empty or "list_date" not in catalog.columns:
        return pd.DataFrame(columns=["style", "sku", "sales", "return_qty", "return_rate"])

    threshold = ref_date - timedelta(days=NEW_PRODUCT_DAYS)
    new_skus = catalog[catalog["list_date"] >= threshold]["sku_root"].dropna().unique().tolist()
    if not new_skus:
        return pd.DataFrame(columns=["style", "sku", "sales", "return_qty", "return_rate"])

    paid = this_orders["raw_paid"]
    if paid.empty:
        return pd.DataFrame(columns=["style", "sku", "sales", "return_qty", "return_rate"])

    new_paid = paid[paid["sku_root"].isin(new_skus)].copy()
    if new_paid.empty:
        return pd.DataFrame(columns=["style", "sku", "sales", "return_qty", "return_rate"])

    style_sales = (
        new_paid.groupby(["style", "sku_root"])
        .agg(sales=("Quantity", "sum")).reset_index()
        .rename(columns={"sku_root": "sku"})
    )
    style_sales = style_sales.merge(this_returns["by_style"], on="style", how="left").fillna({"return_qty": 0})
    style_sales["return_rate"] = style_sales.apply(
        lambda r: r["return_qty"] / r["sales"] if r["sales"] > 0 else 0, axis=1
    )
    return style_sales.sort_values("sales", ascending=False).head(10)[
        ["style", "sku", "sales", "return_qty", "return_rate"]
    ]


def build_supplier_analysis(this_orders: dict, this_returns: dict,
                             catalog: pd.DataFrame) -> pd.DataFrame:
    """
    按厂家聚合所有在售款的整体退货率（全新逻辑）。

    输出列：supplier | active_count | sold_count | total_sales | total_return | return_rate | top_problem
    """
    if catalog.empty:
        return pd.DataFrame(columns=["supplier", "active_count", "sold_count",
                                      "total_sales", "total_return", "return_rate", "top_problem"])

    # 在售款（图册里 is_active=True）
    active = catalog[catalog["is_active"] == True].copy()
    active["supplier"] = active["supplier"].replace("", "未知").fillna("未知")

    # 各供应商在售款数
    supplier_active = active.groupby("supplier").size().rename("active_count").reset_index()

    # 销量数据
    paid = this_orders["raw_paid"]
    if not paid.empty:
        # 关联 supplier
        sku_to_supplier = active.set_index("sku_root")["supplier"].to_dict()
        paid = paid.copy()
        paid["supplier"] = paid["sku_root"].map(sku_to_supplier).fillna("未知")

        # 每个款的销量
        style_sales = paid.groupby(["supplier", "style", "sku_root"]).agg(
            sales=("Quantity", "sum")
        ).reset_index()
    else:
        style_sales = pd.DataFrame(columns=["supplier", "style", "sku_root", "sales"])

    # 关联退货
    style_sales = style_sales.merge(this_returns["by_style"], on="style", how="left").fillna({"return_qty": 0})

    # 按供应商聚合
    supplier_metrics = style_sales.groupby("supplier").agg(
        sold_count=("style", lambda s: s.nunique()),  # 有销量的款数
        total_sales=("sales", "sum"),
        total_return=("return_qty", "sum"),
    ).reset_index()

    # 找每个供应商的代表问题款（退货率最高的款，需 sales >= 5）
    style_sales["return_rate"] = style_sales.apply(
        lambda r: r["return_qty"] / r["sales"] if r["sales"] > 0 else 0, axis=1
    )
    problem_per_supplier = {}
    for sup, sub in style_sales[style_sales["sales"] >= 5].groupby("supplier"):
        top = sub.sort_values("return_rate", ascending=False).head(2)
        problem_per_supplier[sup] = ", ".join([
            f"{r['style']}({r['return_rate']:.0%})" for _, r in top.iterrows()
            if r["return_rate"] > 0
        ])

    supplier_metrics["top_problem"] = supplier_metrics["supplier"].map(problem_per_supplier).fillna("—")

    # 合并：在售款数 + 销量数据
    result = supplier_active.merge(supplier_metrics, on="supplier", how="left").fillna({
        "sold_count": 0, "total_sales": 0, "total_return": 0, "top_problem": "—"
    })

    # 退货率
    result["return_rate"] = result.apply(
        lambda r: r["total_return"] / r["total_sales"] if r["total_sales"] > 0 else 0, axis=1
    )

    # 按总销量排序（销量大的供应商在前）
    result = result.sort_values("total_sales", ascending=False).reset_index(drop=True)
    result["sold_count"] = result["sold_count"].astype(int)
    result["total_sales"] = result["total_sales"].astype(int)
    result["total_return"] = result["total_return"].astype(int)
    result["active_count"] = result["active_count"].astype(int)
    return result[["supplier", "active_count", "sold_count", "total_sales",
                   "total_return", "return_rate", "top_problem"]]


# ╔════════════════════════════════════════════════════════════════════╗
# ║   第 8 部分：客诉中台 KPI 分析                                     ║
# ╚════════════════════════════════════════════════════════════════════╝

def build_response_time_analysis(last_returns: dict, this_returns: dict) -> dict:
    """
    客诉响应时效（退款处理周期）分析。
    用 Refund Time - Time Requested 计算每单处理小时数，分桶统计。
    """
    def _calc(returns_dict):
        df = returns_dict.get("raw_actual", pd.DataFrame())
        if df.empty:
            return {"buckets": {b[0]: 0 for b in RESPONSE_TIME_BUCKETS}, "median": 0, "mean": 0, "n": 0}

        # 计算处理时长
        time_col = _find_col(df, RETURN_COL_CANDIDATES["time_requested"])
        refund_col = _find_col(df, RETURN_COL_CANDIDATES["refund_time"])
        if not time_col or not refund_col:
            return {"buckets": {b[0]: 0 for b in RESPONSE_TIME_BUCKETS}, "median": 0, "mean": 0, "n": 0}

        df = df.copy()
        t_req = pd.to_datetime(_clean_tab_str(df[time_col]),
                                format="%d/%m/%Y %H:%M:%S", errors="coerce")
        t_refund = pd.to_datetime(_clean_tab_str(df[refund_col]),
                                   format="%d/%m/%Y %H:%M:%S", errors="coerce")
        df["__hours"] = (t_refund - t_req).dt.total_seconds() / 3600
        df = df[df["__hours"].notna() & (df["__hours"] >= 0)]

        if df.empty:
            return {"buckets": {b[0]: 0 for b in RESPONSE_TIME_BUCKETS}, "median": 0, "mean": 0, "n": 0}

        buckets = {}
        for label, lo, hi in RESPONSE_TIME_BUCKETS:
            buckets[label] = int(((df["__hours"] >= lo) & (df["__hours"] < hi)).sum())

        return {
            "buckets": buckets,
            "median": float(df["__hours"].median()),
            "mean": float(df["__hours"].mean()),
            "n": len(df),
        }

    return {"last": _calc(last_returns), "this": _calc(this_returns)}


def build_repeat_buyers(this_returns: dict, top_n: int = 10) -> pd.DataFrame:
    """重复退货买家 TOP 10（含主要原因 + 建议操作）"""
    df = this_returns.get("raw_all", pd.DataFrame())
    if df.empty:
        return pd.DataFrame(columns=["buyer", "return_count", "amount", "main_reason", "suggestion"])

    buyer_col = _find_col(df, RETURN_COL_CANDIDATES["buyer_username"])
    qty_col = _find_col(df, RETURN_COL_CANDIDATES["qty"])
    unit_price_col = _find_col(df, RETURN_COL_CANDIDATES["unit_price"])
    reason_col = _find_col(df, RETURN_COL_CANDIDATES["reason"])

    if not buyer_col or not qty_col:
        return pd.DataFrame(columns=["buyer", "return_count", "amount", "main_reason", "suggestion"])

    df = df.copy()
    if unit_price_col:
        df[unit_price_col] = pd.to_numeric(df[unit_price_col], errors="coerce").fillna(0)
        df["__amt"] = df[unit_price_col] * df[qty_col]
    else:
        df["__amt"] = 0

    grouped = df.groupby(buyer_col).agg(
        return_count=(qty_col, "count"),
        amount=("__amt", "sum"),
    ).reset_index().rename(columns={buyer_col: "buyer"})

    # 主要原因
    if reason_col:
        main_reason = df.groupby(buyer_col)[reason_col].agg(
            lambda s: s.mode().iloc[0] if not s.mode().empty else ""
        ).reset_index().rename(columns={buyer_col: "buyer", reason_col: "main_reason"})
        grouped = grouped.merge(main_reason, on="buyer", how="left")
    else:
        grouped["main_reason"] = ""

    grouped = grouped[grouped["return_count"] >= 2].sort_values("return_count", ascending=False).head(top_n)

    # 建议操作
    def _suggest(n):
        if n >= 8: return "🚫 建议拉黑"
        if n >= 5: return "⚠️ 重点监控"
        if n >= 3: return "👀 关注"
        return "—"
    grouped["suggestion"] = grouped["return_count"].apply(_suggest)

    grouped["amount"] = grouped["amount"].round(2)
    grouped["return_count"] = grouped["return_count"].astype(int)
    return grouped[["buyer", "return_count", "amount", "main_reason", "suggestion"]]


def build_keyword_analysis(this_returns: dict, top_n: int = 10) -> pd.DataFrame:
    """从 Buyer Note 抽取高频差评关键词。"""
    df = this_returns.get("raw_all", pd.DataFrame())
    if df.empty:
        return pd.DataFrame(columns=["keyword", "count", "ratio", "top_styles"])

    note_col = _find_col(df, RETURN_COL_CANDIDATES["buyer_note"])
    if not note_col:
        return pd.DataFrame(columns=["keyword", "count", "ratio", "top_styles"])

    df = df.copy()
    df["__note"] = df[note_col].fillna("").astype(str).str.lower()
    notes_with_text = df[df["__note"].str.strip() != ""].copy()
    if notes_with_text.empty:
        return pd.DataFrame(columns=["keyword", "count", "ratio", "top_styles"])

    total_notes = len(notes_with_text)

    # 1. 优先匹配重要短语
    phrase_counts = Counter()
    phrase_styles = {}
    for phrase in IMPORTANT_PHRASES:
        mask = notes_with_text["__note"].str.contains(re.escape(phrase), regex=True)
        cnt = int(mask.sum())
        if cnt > 0:
            phrase_counts[phrase] = cnt
            top_s = notes_with_text[mask].groupby("style").size().sort_values(ascending=False).head(2)
            phrase_styles[phrase] = "; ".join([f"{s}({c})" for s, c in top_s.items() if s])

    # 2. 单词级（去停用词）
    word_counts = Counter()
    word_styles = {}
    for _, row in notes_with_text.iterrows():
        text = row["__note"]
        words = re.findall(r"[a-z]{3,}", text)
        words = [w for w in words if w not in STOP_WORDS]
        seen = set()
        for w in words:
            if w in seen:
                continue
            seen.add(w)
            word_counts[w] += 1
            word_styles.setdefault(w, []).append(row["style"])

    # 合并：短语优先 + 排名前 N 的单词补充
    combined = []
    for phrase, cnt in phrase_counts.most_common():
        combined.append({
            "keyword": phrase,
            "count": cnt,
            "ratio": cnt / total_notes,
            "top_styles": phrase_styles.get(phrase, ""),
        })

    # 单词补充（避免和短语重复）
    used_words = set()
    for phrase in phrase_counts:
        for w in phrase.split():
            used_words.add(w)

    for w, cnt in word_counts.most_common(50):
        if w in used_words:
            continue
        if cnt < 3:
            break
        styles_top = Counter(word_styles[w]).most_common(2)
        combined.append({
            "keyword": w,
            "count": cnt,
            "ratio": cnt / total_notes,
            "top_styles": "; ".join([f"{s}({c})" for s, c in styles_top if s]),
        })

    result = pd.DataFrame(combined).sort_values("count", ascending=False).head(top_n)
    return result.reset_index(drop=True)


def build_missing_package_analysis(last_returns: dict, this_returns: dict,
                                    catalog: pd.DataFrame) -> dict:
    """包裹丢失分析"""
    def _count_missing(returns_dict):
        df = returns_dict.get("raw_all", pd.DataFrame())
        if df.empty:
            return 0, 0, pd.DataFrame()
        reason_col = _find_col(df, RETURN_COL_CANDIDATES["reason"])
        qty_col = _find_col(df, RETURN_COL_CANDIDATES["qty"])
        if not reason_col or not qty_col:
            return 0, 0, pd.DataFrame()

        missing = df[df[reason_col] == "Missing package"]
        total_returns = int(df[qty_col].sum())
        missing_qty = int(missing[qty_col].sum())
        return missing_qty, total_returns, missing

    last_qty, last_total, _ = _count_missing(last_returns)
    this_qty, this_total, this_missing = _count_missing(this_returns)

    # 高发款式
    top_styles = pd.DataFrame(columns=["style", "sku", "missing_count", "ratio_to_sales"])
    if not this_missing.empty:
        qty_col = _find_col(this_missing, RETURN_COL_CANDIDATES["qty"])
        ts = this_missing.groupby(["style", "sku_root"]).agg(
            missing_count=(qty_col, "sum")
        ).reset_index().sort_values("missing_count", ascending=False).head(5)
        ts["sku"] = ts["sku_root"]
        ts["ratio_to_sales"] = 0
        top_styles = ts[["style", "sku", "missing_count", "ratio_to_sales"]]

    return {
        "last_qty": last_qty,
        "this_qty": this_qty,
        "last_ratio": last_qty / last_total if last_total else 0,
        "this_ratio": this_qty / this_total if this_total else 0,
        "wow": _wow(this_qty, last_qty),
        "top_styles": top_styles,
    }


def build_lifecycle_analysis(this_orders: dict, this_returns: dict,
                              catalog: pd.DataFrame, ref_date: datetime) -> pd.DataFrame:
    """按上架时长分组看退货率"""
    if catalog.empty or this_orders["raw_paid"].empty:
        return pd.DataFrame(columns=["lifecycle", "active_count", "sales", "return_qty", "return_rate"])

    cat = catalog[catalog["is_active"] == True].copy()
    cat["__days"] = (ref_date - cat["list_date"]).dt.days

    paid = this_orders["raw_paid"]
    by_style = this_returns["by_style"]

    # 计算每个款的销量和退货
    style_data = paid.groupby(["style", "sku_root"]).agg(sales=("Quantity", "sum")).reset_index()
    style_data = style_data.merge(by_style, on="style", how="left").fillna({"return_qty": 0})

    # 关联 lifecycle
    sku_to_days = cat.set_index("sku_root")["__days"].to_dict()
    style_data["__days"] = style_data["sku_root"].map(sku_to_days)
    style_data = style_data[style_data["__days"].notna()].copy()

    def _bucket(d):
        for label, lo, hi in LIFECYCLE_BUCKETS:
            if lo <= d < hi:
                return label
        return LIFECYCLE_BUCKETS[-1][0]

    style_data["lifecycle"] = style_data["__days"].apply(_bucket)

    # 各 lifecycle 的在售款数（从 catalog 算，更准）
    cat["lifecycle"] = cat["__days"].apply(lambda d: _bucket(d) if pd.notna(d) else None)
    active_per_lc = cat.groupby("lifecycle").size().rename("active_count").reset_index()

    # 销量&退货
    sales_per_lc = style_data.groupby("lifecycle").agg(
        sales=("sales", "sum"),
        return_qty=("return_qty", "sum"),
    ).reset_index()

    # 合并
    result = active_per_lc.merge(sales_per_lc, on="lifecycle", how="outer").fillna(0)
    result["return_rate"] = result.apply(
        lambda r: r["return_qty"] / r["sales"] if r["sales"] > 0 else 0, axis=1
    )

    # 排序按预定顺序
    order = [b[0] for b in LIFECYCLE_BUCKETS]
    result["__order"] = result["lifecycle"].apply(lambda x: order.index(x) if x in order else 99)
    result = result.sort_values("__order").drop(columns="__order").reset_index(drop=True)

    for col in ["active_count", "sales", "return_qty"]:
        result[col] = result[col].astype(int)

    return result[["lifecycle", "active_count", "sales", "return_qty", "return_rate"]]


# ╔════════════════════════════════════════════════════════════════════╗
# ║   第 9 部分：写入 Excel 模板                                       ║
# ╚════════════════════════════════════════════════════════════════════╝

def _write_template(template_bytes, last_label, this_label,
                    last_o, this_o, last_r, this_r,
                    top10_sales, top10_return,
                    new_products_top10, supplier_analysis,
                    response_time, repeat_buyers, keywords,
                    missing_package, lifecycle,
                    traffic_metrics) -> bytes:
    """把所有指标写到模板的副本里。"""
    wb = load_workbook(io.BytesIO(template_bytes))

    if BASE_SHEET in wb.sheetnames:
        base_ws = wb[BASE_SHEET]
    else:
        base_ws = wb[wb.sheetnames[-1]]

    ws = wb.copy_worksheet(base_ws)
    safe_name = this_label.replace(":", "_").replace("/", "-")[:31]
    ws.title = safe_name

    # ---------------- B2 标题 ----------------
    ws["B2"] = f"NailVesta 中台运营周报  |  Week：{this_label}"

    # ---------------- 核心指标区 ----------------
    metrics_rows_last = {6: last_o["orders"], 7: last_o["sku_sold"], 8: last_o["attach_rate"],
                          9: last_o["asp"], 10: last_o["aov"], 15: last_o["gmv"]}
    metrics_rows_this = {6: this_o["orders"], 7: this_o["sku_sold"], 8: this_o["attach_rate"],
                          9: this_o["asp"], 10: this_o["aov"], 15: this_o["gmv"]}
    for row, last_v in metrics_rows_last.items():
        ws.cell(row=row, column=3, value=last_v)
        ws.cell(row=row, column=6, value=metrics_rows_this[row])
        ws.cell(row=row, column=7, value=f"=IFERROR(F{row}/C{row}-1,\"\")")
    ws["D6"] = last_o["influencer_orders"]
    ws["E6"] = this_o["influencer_orders"]
    ws["D7"] = ""
    ws["E7"] = ""

    # 流量指标
    flow_map = {11: ("ctr_last", "ctr_this"), 12: ("cvr_last", "cvr_this"),
                13: ("cart_conv_last", "cart_conv_this"), 14: ("atc_last", "atc_this")}
    for row, (lk, tk) in flow_map.items():
        lv = _parse_pct(traffic_metrics.get(lk, ""))
        tv = _parse_pct(traffic_metrics.get(tk, ""))
        if lv is not None:
            ws.cell(row=row, column=3, value=lv).number_format = "0.00%"
        if tv is not None:
            ws.cell(row=row, column=6, value=tv).number_format = "0.00%"

    # ---------------- 退换货总览 ----------------
    rr_last = last_r["total_qty"] / last_o["sku_sold"] if last_o["sku_sold"] else 0
    rr_this = this_r["total_qty"] / this_o["sku_sold"] if this_o["sku_sold"] else 0
    per_last = last_r["total_amount"] / last_r["total_qty"] if last_r["total_qty"] else 0
    per_this = this_r["total_amount"] / this_r["total_qty"] if this_r["total_qty"] else 0
    return_rows = {
        6:  (last_r["total_qty"],     this_r["total_qty"]),
        7:  (rr_last,                 rr_this),
        8:  (last_r["total_amount"],  this_r["total_amount"]),
        9:  (per_last,                per_this),
        10: (len(last_r["by_style"]), len(this_r["by_style"])),
    }
    for row, (lv, tv) in return_rows.items():
        ws.cell(row=row, column=13, value=lv)
        ws.cell(row=row, column=14, value=tv)
        ws.cell(row=row, column=15, value=f"=IFERROR(N{row}/M{row}-1,\"\")")
    ws["M7"].number_format = "0.00%"
    ws["N7"].number_format = "0.00%"

    # ---------------- 订单内 SKU 数量结构 ----------------
    bucket_rows = {"1": 18, "2": 19, "3": 20, "4": 21, "4+": 22}
    last_bucket = last_o["bucket_struct"].set_index("bucket")["order_count"].to_dict()
    this_bucket = this_o["bucket_struct"].set_index("bucket")["order_count"].to_dict()
    last_total = sum(last_bucket.values()) or 1
    this_total = sum(this_bucket.values()) or 1
    for b, row in bucket_rows.items():
        lo = last_bucket.get(b, 0); to = this_bucket.get(b, 0)
        ws.cell(row=row, column=3, value=lo)
        ws.cell(row=row, column=4, value=lo / last_total).number_format = "0.00%"
        ws.cell(row=row, column=5, value=round(last_o["bucket_aov"].get(b, 0), 2))
        ws.cell(row=row, column=6, value=to)
        ws.cell(row=row, column=7, value=to / this_total).number_format = "0.00%"
        ws.cell(row=row, column=8, value=round(this_o["bucket_aov"].get(b, 0), 2))

    # ---------------- 退货原因 ----------------
    reason_rows_map = {
        "No longer needed": 18, "Missing package": 19, "Wrong item was sent": 20,
        "Item doesn't match description": 21, "Defective item": 22,
        "Product wouldn't arrive on time": 23,
        "Congrats on meeting your refundable sample criteria!": 24,
        "Missing items": 25, "Damaged item or packaging": 26,
    }
    last_reason = last_r["by_reason"].set_index("reason")["qty"].to_dict() if not last_r["by_reason"].empty else {}
    this_reason = this_r["by_reason"].set_index("reason")["qty"].to_dict() if not this_r["by_reason"].empty else {}
    last_reason_total = sum(last_reason.values()) or 1
    this_reason_total = sum(this_reason.values()) or 1

    def _match_reason(target, reason_dict):
        target_l = target.lower()
        for k, v in reason_dict.items():
            if k and (target_l in str(k).lower() or str(k).lower() in target_l):
                return v
        return 0

    for reason, row in reason_rows_map.items():
        lq = _match_reason(reason, last_reason)
        tq = _match_reason(reason, this_reason)
        ws.cell(row=row, column=12, value=reason)
        ws.cell(row=row, column=13, value=lq)
        ws.cell(row=row, column=14, value=lq / last_reason_total).number_format = "0.00%"
        ws.cell(row=row, column=15, value=tq)
        ws.cell(row=row, column=16, value=tq / this_reason_total).number_format = "0.00%"
        top_style = ""
        for k, v in this_r["by_reason_top_style"].items():
            if reason.lower() in str(k).lower() or str(k).lower() in reason.lower():
                top_style = v; break
        ws.cell(row=row, column=17, value=top_style)

    # ---------------- 销量 TOP 10 (B30:F39) ----------------
    for i, (_, rd) in enumerate(top10_sales.iterrows()):
        r = 30 + i
        ws.cell(row=r, column=2, value=rd["style"])
        ws.cell(row=r, column=3, value=int(rd["sales"]))
        ws.cell(row=r, column=4, value=int(rd["return_qty"]))
        ws.cell(row=r, column=5, value=rd["return_rate"]).number_format = "0.00%"
        ws.cell(row=r, column=6, value="")
    for i in range(len(top10_sales), 10):
        r = 30 + i
        for col in range(2, 7):
            ws.cell(row=r, column=col, value="")

    # ---------------- 退货率 TOP 10 (L30:P39) ----------------
    for i, (_, rd) in enumerate(top10_return.iterrows()):
        r = 30 + i
        ws.cell(row=r, column=12, value=rd["style"])
        ws.cell(row=r, column=13, value=int(rd["sales"]))
        ws.cell(row=r, column=14, value=int(rd["return_qty"]))
        ws.cell(row=r, column=15, value=rd["return_rate"]).number_format = "0.00%"
        ws.cell(row=r, column=16, value="")
    for i in range(len(top10_return), 10):
        r = 30 + i
        for col in range(12, 17):
            ws.cell(row=r, column=col, value="")

    # ---------------- 🌟 近30天新品 (B43:J52) ----------------
    for i in range(10):
        r = 43 + i
        if i < len(new_products_top10):
            rd = new_products_top10.iloc[i]
            ws.cell(row=r, column=2, value=rd["style"])
            ws.cell(row=r, column=3, value=rd["sku"])
            ws.cell(row=r, column=6, value=int(rd["sales"]))
            ws.cell(row=r, column=7, value=int(rd["return_qty"]))
            ws.cell(row=r, column=10, value=rd["return_rate"]).number_format = "0.00%"
        else:
            for col in [2, 3, 6, 7, 10]:
                ws.cell(row=r, column=col, value="")

    # ---------------- 🚨 供应商分析 (L43:R52) ----------------
    # 列：L=供应商 M=在售款数 N=有销量款数 O=总销量 P=总退货 Q=整体退货率 R=代表问题款
    for i in range(10):
        r = 43 + i
        if i < len(supplier_analysis):
            rd = supplier_analysis.iloc[i]
            ws.cell(row=r, column=12, value=rd["supplier"])
            ws.cell(row=r, column=13, value=int(rd["active_count"]))
            ws.cell(row=r, column=14, value=int(rd["sold_count"]))
            ws.cell(row=r, column=15, value=int(rd["total_sales"]))
            ws.cell(row=r, column=16, value=int(rd["total_return"]))
            ws.cell(row=r, column=17, value=rd["return_rate"]).number_format = "0.00%"
            ws.cell(row=r, column=18, value=rd["top_problem"])
        else:
            for col in range(12, 19):
                ws.cell(row=r, column=col, value="")

    # ╔══════════════════════════════════════════════════════════════╗
    # ║   新增的 5 个客诉中台 KPI 区域                              ║
    # ╚══════════════════════════════════════════════════════════════╝

    # ---------------- ⏱️ 客诉响应时效 (B71:G78) ----------------
    last_rt = response_time["last"]; this_rt = response_time["this"]
    last_rt_total = sum(last_rt["buckets"].values()) or 1
    this_rt_total = sum(this_rt["buckets"].values()) or 1
    for i, (label, _, _) in enumerate(RESPONSE_TIME_BUCKETS):
        r = 73 + i
        lq = last_rt["buckets"][label]
        tq = this_rt["buckets"][label]
        ws.cell(row=r, column=3, value=lq)
        ws.cell(row=r, column=4, value=lq / last_rt_total).number_format = "0.00%"
        ws.cell(row=r, column=5, value=tq)
        ws.cell(row=r, column=6, value=tq / this_rt_total).number_format = "0.00%"
        ws.cell(row=r, column=7, value=f"=IFERROR(E{r}/C{r}-1,\"\")")

    # 中位数 / 平均数（row 78）
    ws.cell(row=78, column=3, value=f"{last_rt['median']:.1f}h / {last_rt['mean']:.1f}h"
            if last_rt["n"] else "—")
    ws.cell(row=78, column=5, value=f"{this_rt['median']:.1f}h / {this_rt['mean']:.1f}h"
            if this_rt["n"] else "—")

    # ---------------- 🔁 重复退货买家 (B82:G91) ----------------
    for i in range(10):
        r = 82 + i
        if i < len(repeat_buyers):
            rd = repeat_buyers.iloc[i]
            ws.cell(row=r, column=3, value=rd["buyer"])
            ws.cell(row=r, column=4, value=int(rd["return_count"]))
            ws.cell(row=r, column=5, value=float(rd["amount"]))
            ws.cell(row=r, column=6, value=str(rd["main_reason"])[:30])
            ws.cell(row=r, column=7, value=rd["suggestion"])
        else:
            for col in range(3, 8):
                ws.cell(row=r, column=col, value="")

    # ---------------- 💬 差评关键词 (B95:G104) ----------------
    for i in range(10):
        r = 95 + i
        if i < len(keywords):
            rd = keywords.iloc[i]
            ws.cell(row=r, column=3, value=rd["keyword"])
            ws.cell(row=r, column=4, value=int(rd["count"]))
            ws.cell(row=r, column=5, value=rd["ratio"]).number_format = "0.00%"
            ws.cell(row=r, column=6, value=str(rd["top_styles"]))
        else:
            for col in range(3, 7):
                ws.cell(row=r, column=col, value="")

    # ---------------- 📦 包裹丢失 (B108:G108, B112:G116) ----------------
    ws.cell(row=108, column=2, value=missing_package["last_qty"])
    ws.cell(row=108, column=3, value=missing_package["this_qty"])
    ws.cell(row=108, column=4, value=missing_package["wow"]).number_format = "0.00%"
    ws.cell(row=108, column=5, value=missing_package["this_ratio"]).number_format = "0.00%"

    for i in range(5):
        r = 112 + i
        if i < len(missing_package["top_styles"]):
            rd = missing_package["top_styles"].iloc[i]
            ws.cell(row=r, column=3, value=rd["style"])
            ws.cell(row=r, column=4, value=rd["sku"])
            ws.cell(row=r, column=5, value=int(rd["missing_count"]))
            ws.cell(row=r, column=6, value="—")  # 占该款销量比，简化
        else:
            for col in range(3, 7):
                ws.cell(row=r, column=col, value="")

    # ---------------- 📈 产品生命周期 (B120:F124) ----------------
    for i in range(4):
        r = 120 + i
        if i < len(lifecycle):
            rd = lifecycle.iloc[i]
            ws.cell(row=r, column=3, value=int(rd["active_count"]))
            ws.cell(row=r, column=4, value=int(rd["sales"]))
            ws.cell(row=r, column=5, value=int(rd["return_qty"]))
            ws.cell(row=r, column=6, value=rd["return_rate"]).number_format = "0.00%"
        else:
            for col in range(3, 7):
                ws.cell(row=r, column=col, value="")

    # 合计行 (row 124)
    if not lifecycle.empty:
        ws.cell(row=124, column=3, value=int(lifecycle["active_count"].sum()))
        ws.cell(row=124, column=4, value=int(lifecycle["sales"].sum()))
        ws.cell(row=124, column=5, value=int(lifecycle["return_qty"].sum()))
        total_rate = (lifecycle["return_qty"].sum() / lifecycle["sales"].sum()
                       if lifecycle["sales"].sum() > 0 else 0)
        ws.cell(row=124, column=6, value=total_rate).number_format = "0.00%"

    # ---------------- 输出 ----------------
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


# ╔════════════════════════════════════════════════════════════════════╗
# ║   第 10 部分：主入口                                               ║
# ╚════════════════════════════════════════════════════════════════════╝

def build_report(orders_file, returns_file, catalog_file,
                 template_bytes: bytes,
                 time_window_mode: str,
                 influencer_rule: str,
                 traffic_metrics: dict,
                 exclude_request_canceled: bool = True,
                 custom_last_label: str = None,
                 custom_this_label: str = None):

    influencer_col = INFLUENCER_RULE_MAP[influencer_rule]

    orders_df  = _read_any(orders_file)
    returns_df = _read_any(returns_file)
    catalog_df = _read_any(catalog_file)

    if orders_df.empty:
        raise ValueError("订单数据为空")

    windows = split_time_windows(orders_df, returns_df, time_window_mode)
    last_label = custom_last_label or windows["last_label"]
    this_label = custom_this_label or windows["this_label"]

    last_o = process_orders(windows["last_orders"], influencer_col)
    this_o = process_orders(windows["this_orders"], influencer_col)
    last_r = process_returns(windows["last_returns"], exclude_request_canceled=exclude_request_canceled)
    this_r = process_returns(windows["this_returns"], exclude_request_canceled=exclude_request_canceled)

    catalog = process_catalog(catalog_df)

    # TOP 10
    top10_sales = this_o["style_sales"].head(10).copy()
    top10_sales = top10_sales.merge(this_r["by_style"], on="style", how="left").fillna({"return_qty": 0})
    top10_sales["return_rate"] = top10_sales.apply(
        lambda r: r["return_qty"] / r["sales"] if r["sales"] > 0 else 0, axis=1
    )
    top10_sales = top10_sales[["style", "sales", "return_qty", "return_rate"]]

    style_full = this_o["style_sales"].merge(this_r["by_style"], on="style", how="left").fillna({"return_qty": 0})
    style_full = style_full[style_full["sales"] >= 10].copy()
    style_full["return_rate"] = style_full.apply(
        lambda r: r["return_qty"] / r["sales"] if r["sales"] > 0 else 0, axis=1
    )
    top10_return = style_full.sort_values("return_rate", ascending=False).head(10)[
        ["style", "sales", "return_qty", "return_rate"]
    ]

    # 新品 + 供应商
    new_products = build_new_products_top10(this_o, this_r, catalog, windows["ref_date"])
    supplier_analysis = build_supplier_analysis(this_o, this_r, catalog)

    # 客诉 KPI
    response_time = build_response_time_analysis(last_r, this_r)
    repeat_buyers = build_repeat_buyers(this_r)
    keywords = build_keyword_analysis(this_r)
    missing_pkg = build_missing_package_analysis(last_r, this_r, catalog)
    lifecycle = build_lifecycle_analysis(this_o, this_r, catalog, windows["ref_date"])

    # 写模板
    output_bytes = _write_template(
        template_bytes=template_bytes,
        last_label=last_label, this_label=this_label,
        last_o=last_o, this_o=this_o, last_r=last_r, this_r=this_r,
        top10_sales=top10_sales, top10_return=top10_return,
        new_products_top10=new_products,
        supplier_analysis=supplier_analysis,
        response_time=response_time,
        repeat_buyers=repeat_buyers,
        keywords=keywords,
        missing_package=missing_pkg,
        lifecycle=lifecycle,
        traffic_metrics=traffic_metrics,
    )

    # 摘要
    rr_last = last_r["total_qty"] / last_o["sku_sold"] if last_o["sku_sold"] else 0
    rr_this = this_r["total_qty"] / this_o["sku_sold"] if this_o["sku_sold"] else 0
    base_keys = ["orders", "sku_sold", "gmv", "aov", "asp", "attach_rate", "influencer_orders"]
    summary = {
        "last": {**{k: last_o[k] for k in base_keys},
                 "return_rate": rr_last,
                 "total_qty": last_r["total_qty"],
                 "requested_qty": last_r["requested_qty"],
                 "request_canceled_qty": last_r["request_canceled_qty"],
                 "label": last_label},
        "this": {**{k: this_o[k] for k in base_keys},
                 "return_rate": rr_this,
                 "total_qty": this_r["total_qty"],
                 "requested_qty": this_r["requested_qty"],
                 "request_canceled_qty": this_r["request_canceled_qty"],
                 "label": this_label},
        "wow": {
            "orders":      _wow(this_o["orders"], last_o["orders"]),
            "gmv":         _wow(this_o["gmv"], last_o["gmv"]),
            "aov":         _wow(this_o["aov"], last_o["aov"]),
            "return_rate": rr_this - rr_last,
        },
        "top10_sales": top10_sales.assign(
            return_rate=top10_sales["return_rate"].apply(lambda x: f"{x:.1%}")
        ),
        "top10_return": top10_return.assign(
            return_rate=top10_return["return_rate"].apply(lambda x: f"{x:.1%}")
        ),
        "new_products": new_products.assign(
            return_rate=new_products["return_rate"].apply(lambda x: f"{x:.1%}")
        ) if not new_products.empty else new_products,
        "supplier_analysis": supplier_analysis.assign(
            return_rate=supplier_analysis["return_rate"].apply(lambda x: f"{x:.2%}")
        ) if not supplier_analysis.empty else supplier_analysis,
        "response_time": response_time,
        "repeat_buyers": repeat_buyers,
        "keywords": keywords.assign(
            ratio=keywords["ratio"].apply(lambda x: f"{x:.1%}")
        ) if not keywords.empty else keywords,
        "missing_package": missing_pkg,
        "lifecycle": lifecycle.assign(
            return_rate=lifecycle["return_rate"].apply(lambda x: f"{x:.2%}")
        ) if not lifecycle.empty else lifecycle,
        "windows": {
            "t_min": windows["t_min"],
            "t_max": windows["t_max"],
            "ref_date": windows["ref_date"],
        },
    }
    return output_bytes, summary


# ╔════════════════════════════════════════════════════════════════════╗
# ║   第 11 部分：Streamlit UI                                         ║
# ╚════════════════════════════════════════════════════════════════════╝

st.set_page_config(page_title="NailVesta 周报生成器 v3", page_icon="💅", layout="wide")
st.title("💅 NailVesta 中台运营周报生成器 v3")
st.caption("28 天数据 + 产品图册 → 全量供应商分析 + 客诉中台 KPI")

with st.sidebar:
    st.header("⚙️ 参数")

    st.subheader("🕐 时间窗模式")
    time_window_mode_label = st.radio(
        "选择对比方式",
        ["28天对半切（前14 vs 后14）", "自然周对比（最近周 vs 上一周）"],
        index=0,
    )
    time_window_mode = "halve" if "对半切" in time_window_mode_label else "weekly"

    st.divider()
    st.subheader("周次标签覆盖（可选）")
    custom_last = st.text_input("上周标签", value="")
    custom_this = st.text_input("本周标签", value="")

    st.divider()
    st.subheader("达人单识别")
    influencer_rule = st.selectbox("判定为达人单的条件", list(INFLUENCER_RULE_MAP.keys()), index=0)

    st.divider()
    st.subheader("退货过滤")
    exclude_request_canceled = st.checkbox(
        "排除「Request Canceled」(买家主动撤销)", value=True,
        help="买家撤销不计入退货指标，但保留在退货原因统计。",
    )

    st.divider()
    st.subheader("可选 - TikTok 流量数据")
    traffic_metrics = {
        "ctr_last":       st.text_input("上周 CTR (%)"),
        "ctr_this":       st.text_input("本周 CTR (%)"),
        "cvr_last":       st.text_input("上周 CVR (%)"),
        "cvr_this":       st.text_input("本周 CVR (%)"),
        "cart_conv_last": st.text_input("上周 Cart Conv (%)"),
        "cart_conv_this": st.text_input("本周 Cart Conv (%)"),
        "atc_last":       st.text_input("上周 Add-To-Cart (%)"),
        "atc_this":       st.text_input("本周 Add-To-Cart (%)"),
    }

st.subheader("📂 上传 3 个文件")
col1, col2, col3 = st.columns(3)
with col1:
    st.markdown("**1. 订单（28天）**")
    orders_file = st.file_uploader("All_order csv/xlsx", type=["csv", "xlsx"], key="orders")
with col2:
    st.markdown("**2. 退货（28天）**")
    returns_file = st.file_uploader("Return_Refund_Orders xlsx", type=["csv", "xlsx"], key="returns")
with col3:
    st.markdown("**3. 产品图册**")
    DEFAULT_CATALOG = Path(__file__).parent / "NailVesta_产品图册.csv"
    catalog_file = st.file_uploader("NailVesta_产品图册（不传用默认）",
                                     type=["csv", "xlsx"], key="catalog")

DEFAULT_TEMPLATE = Path(__file__).parent / "NailVesta_中台运营周报_模板.xlsx"
with st.expander("（可选）自定义周报模板"):
    template_file = st.file_uploader("上传自定义模板", type=["xlsx"], key="template")

st.divider()

if st.button("🚀 生成周报", type="primary", use_container_width=True):
    if not orders_file:
        st.error("⚠️ 请上传订单文件")
        st.stop()

    if 'template_file' in dir() and template_file is not None:
        template_bytes = template_file.read()
    else:
        if not DEFAULT_TEMPLATE.exists():
            st.error(f"⚠️ 默认模板不存在：{DEFAULT_TEMPLATE}")
            st.stop()
        template_bytes = DEFAULT_TEMPLATE.read_bytes()

    if catalog_file is None:
        if DEFAULT_CATALOG.exists():
            catalog_bytes = DEFAULT_CATALOG.read_bytes()
            catalog_file = io.BytesIO(catalog_bytes)
            catalog_file.name = DEFAULT_CATALOG.name
        else:
            st.warning("⚠️ 未上传图册且无默认图册，无法识别供应商和新品")

    with st.spinner("正在跑数据... ☕"):
        try:
            output_bytes, summary = build_report(
                orders_file=orders_file,
                returns_file=returns_file,
                catalog_file=catalog_file,
                template_bytes=template_bytes,
                time_window_mode=time_window_mode,
                influencer_rule=influencer_rule,
                exclude_request_canceled=exclude_request_canceled,
                traffic_metrics=traffic_metrics,
                custom_last_label=custom_last or None,
                custom_this_label=custom_this or None,
            )
        except Exception as e:
            import traceback
            st.error(f"❌ 生成失败：{e}")
            st.code(traceback.format_exc())
            st.stop()

    st.success("✅ 周报已生成！")

    w = summary["windows"]
    st.info(
        f"📅 数据范围：{w['t_min'].strftime('%Y-%m-%d')} ~ {w['t_max'].strftime('%Y-%m-%d')}  "
        f"|  上周：**{summary['last']['label']}**  vs  本周：**{summary['this']['label']}**"
    )

    # 核心看板
    st.subheader("📊 核心指标")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("订单数", f"{summary['this']['orders']:,}", f"{summary['wow']['orders']:+.1%}")
    c2.metric("总GMV ($)", f"{summary['this']['gmv']:,.0f}", f"{summary['wow']['gmv']:+.1%}")
    c3.metric("AOV ($)", f"{summary['this']['aov']:.2f}", f"{summary['wow']['aov']:+.1%}")
    c4.metric("退货率", f"{summary['this']['return_rate']:.2%}",
              f"{summary['wow']['return_rate']:+.1%}", delta_color="inverse")

    c5, c6, c7, c8 = st.columns(4)
    c5.metric("达人单（上周）", f"{summary['last']['influencer_orders']:,}")
    c6.metric("达人单（本周）", f"{summary['this']['influencer_orders']:,}")
    c7.metric("SKU 销量", f"{summary['this']['sku_sold']:,}")
    c8.metric("连带率", f"{summary['this']['attach_rate']:.2f}")

    # 退货拆解
    st.markdown("**📦 退货拆解** —— 实际退货 = 申请总量 − 买家撤销")
    c9, c10, c11, c12 = st.columns(4)
    c9.metric("申请退货总量", f"{summary['this']['requested_qty']:,}")
    c10.metric("买家撤销", f"{summary['this']['request_canceled_qty']:,}")
    c11.metric("实际退货", f"{summary['this']['total_qty']:,}")
    cancel_rate = (summary['this']['request_canceled_qty'] / summary['this']['requested_qty']
                   if summary['this']['requested_qty'] else 0)
    c12.metric("买家撤销率", f"{cancel_rate:.1%}",
               help="数字越高 = 客服安抚越成功 ✨")

    # TOP 10
    st.subheader("🏆 销量 TOP 10 款式")
    st.dataframe(summary["top10_sales"], use_container_width=True, hide_index=True)
    st.subheader("⚠️ 退货率 TOP 10 款式（销量≥10）")
    st.dataframe(summary["top10_return"], use_container_width=True, hide_index=True)
    st.subheader("🌟 近30天新品 TOP 10")
    if not summary["new_products"].empty:
        st.dataframe(summary["new_products"], use_container_width=True, hide_index=True)
    else:
        st.caption("（暂无新品数据）")

    # 供应商分析（升级版）
    st.subheader("🏭 供应商整体分析（按厂家聚合所有在售款）")
    if not summary["supplier_analysis"].empty:
        st.dataframe(summary["supplier_analysis"], use_container_width=True, hide_index=True)
    else:
        st.caption("（图册未提供供应商字段或无销量）")

    # ============= 客诉中台 KPI =============
    st.divider()
    st.header("👥 客诉中台 KPI 监控")

    # 1. 响应时效
    st.subheader("⏱️ 客诉响应时效")
    rt = summary["response_time"]
    rta, rtb, rtc, rtd = st.columns(4)
    rta.metric("处理中位数（本周）",
               f"{rt['this']['median']:.1f} 小时" if rt['this']['n'] else "—")
    rtb.metric("处理中位数（上周）",
               f"{rt['last']['median']:.1f} 小时" if rt['last']['n'] else "—")
    rtc.metric("可计算单数（本周）", f"{rt['this']['n']}")
    fast_pct_this = (rt['this']['buckets'].get('< 6 小时', 0) +
                     rt['this']['buckets'].get('6-24 小时', 0))
    fast_pct_last = (rt['last']['buckets'].get('< 6 小时', 0) +
                     rt['last']['buckets'].get('6-24 小时', 0))
    n_this = sum(rt['this']['buckets'].values()) or 1
    n_last = sum(rt['last']['buckets'].values()) or 1
    rtd.metric("24小时内处理率",
               f"{fast_pct_this/n_this:.1%}",
               f"{(fast_pct_this/n_this) - (fast_pct_last/n_last):+.1%}")

    bucket_df = pd.DataFrame([
        {"处理时长": k, "上周": rt['last']['buckets'][k], "本周": rt['this']['buckets'][k]}
        for k, _, _ in RESPONSE_TIME_BUCKETS
    ])
    st.dataframe(bucket_df, use_container_width=True, hide_index=True)

    # 2. 重复退货买家
    st.subheader("🔁 重复退货买家 TOP 10")
    if not summary["repeat_buyers"].empty:
        st.dataframe(summary["repeat_buyers"], use_container_width=True, hide_index=True)
    else:
        st.caption("（本周无重复退货）")

    # 3. 关键词
    st.subheader("💬 差评关键词 TOP 10（来自 Buyer Note）")
    if not summary["keywords"].empty:
        st.dataframe(summary["keywords"], use_container_width=True, hide_index=True)
    else:
        st.caption("（本周 Buyer Note 数据不足）")

    # 4. 包裹丢失
    st.subheader("📦 包裹丢失分析")
    mp = summary["missing_package"]
    mpa, mpb, mpc, mpd = st.columns(4)
    mpa.metric("上周丢包数", f"{mp['last_qty']}")
    mpb.metric("本周丢包数", f"{mp['this_qty']}",
               f"{mp['wow']:+.1%}", delta_color="inverse")
    mpc.metric("占总退货比", f"{mp['this_ratio']:.1%}")
    mpd.metric("上周占比", f"{mp['last_ratio']:.1%}")
    if not mp["top_styles"].empty:
        st.markdown("**高发丢包款式 TOP 5**")
        st.dataframe(mp["top_styles"][["style", "sku", "missing_count"]],
                     use_container_width=True, hide_index=True)

    # 5. 生命周期
    st.subheader("📈 产品生命周期退货率")
    if not summary["lifecycle"].empty:
        st.dataframe(summary["lifecycle"], use_container_width=True, hide_index=True)

    # 下载
    st.divider()
    fname = f"NailVesta_周报_{summary['this']['label'].replace('/', '-')}.xlsx"
    st.download_button(
        "📥 下载完整周报 Excel",
        data=output_bytes,
        file_name=fname,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )

else:
    st.info("👆 上传订单 + 退货 + 产品图册后点「生成周报」就行啦~")

    with st.expander("📖 v3 新功能"):
        st.markdown(
            """
            **客诉中台 KPI（新增）**

            1. **⏱️ 客诉响应时效** — 退款处理周期分桶（< 6小时 / 6-24小时 / 1-3天 / 3-7天 / >7天）
            2. **🔁 重复退货买家 TOP 10** — 高频退货者（≥3次自动标记关注，≥5重点监控，≥8建议拉黑）
            3. **💬 差评关键词 TOP 10** — 从 Buyer Note 抽取，含重要短语（"wrong size", "doesn't fit" 等）
            4. **📦 包裹丢失分析** — Missing package 的趋势 + 高发款式
            5. **📈 产品生命周期退货率** — 新品/成长/成熟/长尾期分组对比

            **供应商分析升级**

            从「只看问题款占比」→ **按厂家聚合所有在售款的整体退货率**：
            - 在售款数 / 有销量款数 / 总销量 / 总退货 / 整体退货率 / 代表问题款
            - 不再被供应商体量误导（PONY 量大不代表退货率高）
            """
        )
